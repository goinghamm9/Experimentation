#!/usr/bin/env python3
"""Build a client-facing proposal as .docx and .md from ONE content module.

Usage:
    python3 docx_builder.py CONTENT.py --docx OUT.docx --md OUT.md

CONTENT.py is a plain Python file that defines:
    COVER    = {"wordmark": "Agency Name", "title": "…: proposal", "lines": ["Prepared for: …", "Prepared by: …", ...]}
    SECTIONS = [{"title": "1. Opening", "blocks": [block, ...]}, ...]
    APPENDIX = {"title": "Appendix", "blocks": [...]}          # optional; starts on a new page
    FOOTER_TEXT = "Agency | City, State | Confidential"       # a page number is appended
    META  = {"author": "Sender Name", "title": "…", "company": "Agency", "created": "2026-09-24"}  # optional
    STYLE = {"font": "Arial", "body_pt": 11, "table_pt": 10, "heading_rgb": "1F3864", "space_after_pt": 2}  # optional

Block types (a block is a tuple):
    ("p", text)                          paragraph
    ("bullets", [text, ...])             bulleted list
    ("h2", text)                         sub-heading
    ("table", header, rows, widths)      table; widths in inches, summing to 6.5 or less
    ("sig", text)                        signature line (bold)

Inline formatting inside any text: [bracketed text] is a placeholder and is highlighted
yellow in the docx (kept as-is in the markdown); **double asterisks** render bold.
Both outputs come from the same data, so they cannot drift. Keep content in the
content module, never in this file, so a rebuild after an edit is one command.
"""
import argparse
import importlib.util
import re
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

TOKEN_RE = re.compile(r"(\*\*[^*]+\*\*|\[[^\[\]]*\])")


def load_content(path):
    spec = importlib.util.spec_from_file_location("proposal_content", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    for name in ("COVER", "SECTIONS", "FOOTER_TEXT"):
        if not hasattr(mod, name):
            raise SystemExit(f"content module lacks {name}")
    return mod


def add_runs(paragraph, text, size=None, bold=False, color=None):
    for part in TOKEN_RE.split(text):
        if not part:
            continue
        is_bold, is_ph = bold, False
        if part.startswith("**") and part.endswith("**"):
            part, is_bold = part[2:-2], True
        elif part.startswith("[") and part.endswith("]"):
            is_ph = True
        run = paragraph.add_run(part)
        run.bold = True if is_bold else None
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
        if is_ph:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return paragraph


def set_style_font(style, name, size, bold=None, color=None):
    style.font.name = name
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_page_field(paragraph):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def setup_document(style, footer_text):
    font = style.get("font", "Arial")
    body_pt = style.get("body_pt", 11)
    table_pt = style.get("table_pt", 10)
    navy = RGBColor.from_string(style.get("heading_rgb", "1F3864"))
    space_after = style.get("space_after_pt", 2)

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(1))
    section.footer_distance = Inches(0.5)

    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, font, body_pt)
    normal.paragraph_format.space_after = Pt(space_after)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    for name, size in (("Heading 1", 14), ("Heading 2", 12)):
        h = styles[name]
        set_style_font(h, font, size, bold=True, color=navy)
        h.paragraph_format.space_before = Pt(5)
        h.paragraph_format.space_after = Pt(3)
        h.paragraph_format.keep_with_next = True
    tbl = styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    tbl.base_style = normal
    set_style_font(tbl, font, table_pt)
    tbl.paragraph_format.space_after = Pt(0)
    tbl.paragraph_format.space_before = Pt(0)
    lb = styles["List Bullet"]
    set_style_font(lb, font, body_pt)
    lb.paragraph_format.space_after = Pt(1)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run(footer_text + " | Page ")
    run.font.size = Pt(9)
    run.font.name = font
    add_page_field(fp)
    return doc, navy


def docx_table(doc, header, rows, widths):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = doc.styles["Table Grid"]
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    mar = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), "80")
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)
    hdr = table.rows[0].cells
    for i, text in enumerate(header):
        p = hdr[i].paragraphs[0]
        p.style = doc.styles["Table Text"]
        add_runs(p, text, bold=True)
        shade_cell(hdr[i], "D9D9D9")
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    tr_pr.append(th)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            p = cells[i].paragraphs[0]
            p.style = doc.styles["Table Text"]
            add_runs(p, text)
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    for i, gc in enumerate(table._tbl.tblGrid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[i] * 1440)))
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(0)
    gap.paragraph_format.line_spacing = Pt(3)
    gap.add_run("").font.size = Pt(2)


def docx_block(doc, block):
    kind = block[0]
    if kind == "p":
        add_runs(doc.add_paragraph(), block[1])
    elif kind == "bullets":
        for item in block[1]:
            add_runs(doc.add_paragraph(style="List Bullet"), item)
    elif kind == "h2":
        add_runs(doc.add_paragraph(style="Heading 2"), block[1])
    elif kind == "table":
        docx_table(doc, block[1], block[2], block[3])
    elif kind == "sig":
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        add_runs(p, block[1], bold=True)
    else:
        raise ValueError(f"unknown block type {kind}")


def fix_app_xml(path, company):
    """python-docx's template says 'Microsoft Macintosh Word' and carries an empty
    <Company/>. Replace both in place; the schema allows exactly one Company element."""
    path = str(path)
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/app.xml":
                t = data.decode("utf-8")
                t = re.sub(r"<Application>[^<]*</Application>", "<Application>Microsoft Office Word</Application>", t)
                t = re.sub(r"<Company\s*/>|<Company>[^<]*</Company>", "", t)
                if "<Manager/>" in t:
                    t = t.replace("<Manager/>", f"<Manager/><Company>{company}</Company>", 1)
                else:
                    t = t.replace("</Properties>", f"<Company>{company}</Company></Properties>")
                assert t.count("<Company") == 1
                data = t.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)


def build_docx(mod, out_path):
    style = getattr(mod, "STYLE", {}) or {}
    meta = getattr(mod, "META", {}) or {}
    doc, navy = setup_document(style, mod.FOOTER_TEXT)
    p = doc.add_paragraph()
    add_runs(p, mod.COVER["wordmark"], size=18, bold=True, color=navy)
    p.paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph()
    add_runs(p, mod.COVER["title"], size=16, bold=True)
    p.paragraph_format.space_after = Pt(8)
    for line in mod.COVER.get("lines", []):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        add_runs(p, line, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    for sec in mod.SECTIONS:
        add_runs(doc.add_paragraph(style="Heading 1"), sec["title"])
        for block in sec["blocks"]:
            docx_block(doc, block)
    appendix = getattr(mod, "APPENDIX", None)
    if appendix:
        h = doc.add_paragraph(style="Heading 1")
        h.paragraph_format.page_break_before = True
        add_runs(h, appendix["title"])
        for block in appendix["blocks"]:
            docx_block(doc, block)
    cp = doc.core_properties
    author = meta.get("author", "")
    cp.author = author
    cp.last_modified_by = author
    cp.comments = ""
    cp.title = meta.get("title", mod.COVER["title"])
    cp.subject = cp.keywords = cp.category = ""
    created = meta.get("created")
    if created:
        dt = datetime.strptime(created, "%Y-%m-%d")
        cp.created = cp.modified = dt
    cp.revision = 1
    doc.save(out_path)
    fix_app_xml(out_path, meta.get("company", mod.COVER["wordmark"]))


def md_cell(text):
    return text.replace("|", "\\|")


def md_block(block):
    kind = block[0]
    if kind == "p":
        return block[1] + "\n"
    if kind == "bullets":
        return "\n".join(f"- {item}" for item in block[1]) + "\n"
    if kind == "h2":
        return f"### {block[1]}\n"
    if kind == "table":
        header, rows = block[1], block[2]
        lines = ["| " + " | ".join(md_cell(h) for h in header) + " |", "|" + "---|" * len(header)]
        lines += ["| " + " | ".join(md_cell(c) for c in row) + " |" for row in rows]
        return "\n".join(lines) + "\n"
    if kind == "sig":
        return f"**{block[1]}**\n"
    raise ValueError(kind)


def build_md(mod, out_path):
    out = [f"**{mod.COVER['wordmark']}**\n", f"# {mod.COVER['title']}\n"]
    out += [line + "  " for line in mod.COVER.get("lines", [])]
    out += ["", f"*{mod.FOOTER_TEXT}*\n"]
    for sec in mod.SECTIONS:
        out.append(f"## {sec['title']}\n")
        out += [md_block(b) for b in sec["blocks"]]
    appendix = getattr(mod, "APPENDIX", None)
    if appendix:
        out += ["\n---\n", f"## {appendix['title']}\n"]
        out += [md_block(b) for b in appendix["blocks"]]
    Path(out_path).write_text("\n".join(out), encoding="utf-8")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("content", help="path to the content module (.py)")
    ap.add_argument("--docx", required=True)
    ap.add_argument("--md", required=True)
    args = ap.parse_args()
    mod = load_content(args.content)
    build_docx(mod, args.docx)
    build_md(mod, args.md)
    print(f"wrote {args.docx}\nwrote {args.md}")


if __name__ == "__main__":
    main()
