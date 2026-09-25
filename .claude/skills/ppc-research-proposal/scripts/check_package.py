#!/usr/bin/env python3
"""Quality gate for a research-to-proposal package. Run it before review and again before handoff.
Defaults fit the concise first proposal (body 1 to 4 pages); pass --max-body-pages for longer Stage 2 documents.

Usage:
    python3 check_package.py OUT_DIR --docx 06_proposal.docx [--xlsx 04_budget_model.xlsx]
        [--appendix-heading Appendix] [--max-body-pages 4] [--min-body-pages 1]
        [--banned-names Samprand,OtherName] [--formula-only-tabs "Coverage"]
        [--render 1,3,6] [--json report.json]

Checks (each prints PASS, FAIL, or NOTE with evidence):
  dashes        no em-dash (U+2014) or en-dash (U+2013) in any .md/.csv/.txt, or inside .docx/.xlsx XML
  banned-names  names that must never reach the client (the consultancy behind the agency, for example)
  hype          hype words and promised-result phrases in the client-facing docx and its .md twin
  placeholders  every run containing "[" in the docx is highlighted yellow; counts placeholders
  metadata      docx core author/lastModifiedBy/created/comments; app.xml Application and a single Company
  pages         converts the docx to a FRESH PDF and finds the page where the appendix heading starts;
                the main body must fall within [min, max] pages; optionally renders pages to PNG
  xlsx          formula counts per tab, numeric literals in formula-only tabs, creator metadata, and a
                recalculation with LibreOffice reporting any error values (#REF!, #NAME?, #DIV/0!, ...)

Why a fresh PDF: a stale PDF from an earlier build once reported six pages when the body had grown to seven.
Exit status is 1 when any check FAILs, so it can guard a commit.
"""
import argparse
import glob
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

HYPE = ["cutting-edge", "unlock", "supercharge", "game-changer", "world-class", "state-of-the-art", "leverage",
        "guarantee", "guaranteed", "will generate", "will deliver", "will increase", "best-in-class", "#1"]
DASHES = ("—", "–")
results = []


def rec(check, status, evidence):
    results.append({"check": check, "status": status, "evidence": evidence})
    print(f"[{status}] {check}: {evidence}")


def xml_texts(path):
    z = zipfile.ZipFile(path)
    return {n: z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml")}


def find_soffice():
    for pattern in ("/root/.claude/skills/**/docx/scripts/office/soffice.py", os.path.expanduser("~/.claude/skills/**/docx/scripts/office/soffice.py")):
        hits = glob.glob(pattern, recursive=True)
        if hits:
            return [sys.executable, hits[0]]
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if exe:
        profile = tempfile.mkdtemp(prefix="lo_profile_")
        return [exe, f"-env:UserInstallation=file://{profile}", "--norestore"]
    return None


def convert(path, fmt, outdir):
    cmd = find_soffice()
    if not cmd:
        return None, "LibreOffice not found"
    proc = subprocess.run(cmd + ["--headless", "--convert-to", fmt, "--outdir", outdir, str(path)],
                          capture_output=True, text=True, timeout=300)
    out = Path(outdir) / (Path(path).stem + "." + fmt.split(":")[0])
    return (out if out.exists() else None), (proc.stdout + proc.stderr)[-300:]


def check_dashes(out_dir):
    hits = []
    for f in sorted(Path(out_dir).iterdir()):
        if f.suffix in (".md", ".csv", ".txt"):
            for i, line in enumerate(f.read_text(encoding="utf-8", errors="ignore").splitlines(), 1):
                if any(d in line for d in DASHES):
                    hits.append(f"{f.name}:{i}")
        elif f.suffix in (".docx", ".xlsx"):
            for n, t in xml_texts(f).items():
                if any(d in t for d in DASHES):
                    hits.append(f"{f.name}:{n}")
    rec("dashes", "FAIL" if hits else "PASS", ", ".join(hits[:20]) if hits else "no em or en dashes in any file")


def client_files(out_dir, docx, xlsx):
    files = []
    if docx:
        files.append(Path(out_dir) / docx)
        md = (Path(out_dir) / docx).with_suffix(".md")
        if md.exists():
            files.append(md)
    if xlsx:
        files.append(Path(out_dir) / xlsx)
    return [f for f in files if f.exists()]


def file_text(f):
    if f.suffix in (".docx", ".xlsx"):
        return "\n".join(xml_texts(f).values())
    return f.read_text(encoding="utf-8", errors="ignore")


def check_names(files, names):
    hits = [f"{f.name}:{n}" for f in files for n in names if n and n.lower() in file_text(f).lower()]
    rec("banned-names", "FAIL" if hits else "PASS", ", ".join(hits) if hits else f"none of {names} in client files")


def check_hype(files):
    hits = []
    for f in files:
        if f.suffix == ".xlsx":
            continue
        t = file_text(f).lower()
        hits += [f"{f.name}:{w}" for w in HYPE if w in t]
    rec("hype", "FAIL" if hits else "PASS", ", ".join(hits) if hits else "no hype words or promised results")


def check_placeholders(docx_path):
    import docx as docxlib
    from docx.enum.text import WD_COLOR_INDEX
    d = docxlib.Document(docx_path)
    runs = [r for p in d.paragraphs for r in p.runs]
    runs += [r for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs]
    ph = [r for r in runs if "[" in r.text]
    bad = [r.text[:40] for r in ph if r.font.highlight_color != WD_COLOR_INDEX.YELLOW]
    rec("placeholders", "FAIL" if bad else "PASS", f"{len(ph)} placeholder runs, {len(bad)} not highlighted {bad[:5]}")
    return len(ph)


def check_metadata(docx_path):
    import docx as docxlib
    cp = docxlib.Document(docx_path).core_properties
    xml = xml_texts(docx_path)
    app = xml.get("docProps/app.xml", "")
    application = re.findall(r"<Application>([^<]*)</Application>", app)
    company = re.findall(r"<Company>([^<]*)</Company>", app)
    n_company = app.count("<Company")
    problems = []
    if not cp.author:
        problems.append("author empty")
    if "python-docx" in xml.get("docProps/core.xml", "") or (cp.comments and "python-docx" in cp.comments):
        problems.append("python-docx in core.xml")
    if cp.created and cp.created.year < 2020:
        problems.append(f"created {cp.created.date()} (template default)")
    if n_company != 1:
        problems.append(f"{n_company} Company elements (schema allows one)")
    ev = f"author={cp.author!r} lastModifiedBy={cp.last_modified_by!r} created={cp.created} app={application} company={company}"
    rec("metadata", "FAIL" if problems else "PASS", ev + ("; " + "; ".join(problems) if problems else ""))


def check_pages(docx_path, heading, lo, hi, render):
    try:
        import pymupdf
    except ImportError:
        rec("pages", "NOTE", "pymupdf not installed (pip install pymupdf); page count skipped")
        return
    tmp = tempfile.mkdtemp(prefix="pages_")
    pdf, log = convert(docx_path, "pdf", tmp)
    if not pdf:
        rec("pages", "FAIL", f"PDF conversion failed: {log}")
        return
    d = pymupdf.open(str(pdf))
    total = len(d)
    body = total
    for i, p in enumerate(d):
        if p.get_text().lstrip().lower().startswith(heading.lower()):
            body = i
            break
    last = [l for l in d[body - 1].get_text().splitlines() if l.strip()] if body else []
    status = "PASS" if lo <= body <= hi else "FAIL"
    rec("pages", status, f"total {total} pages, appendix heading on page {body + 1 if body < total else 'n/a'}, main body {body} pages (limit {lo} to {hi}); last body page has {len(last)} text lines")
    for n in render:
        if 1 <= n <= total:
            png = Path(tmp) / f"page{n}.png"
            d[n - 1].get_pixmap(dpi=70).save(str(png))
            print(f"      rendered {png}")
    print(f"      pdf: {pdf}")


def check_xlsx(xlsx_path, formula_tabs):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path)
    summary = []
    literal_hits = []
    for name in wb.sheetnames:
        ws = wb[name]
        cells = [c for row in ws.iter_rows() for c in row if c.value is not None]
        formulas = sum(1 for c in cells if isinstance(c.value, str) and c.value.startswith("="))
        literals = [c.coordinate for c in cells if isinstance(c.value, (int, float))]
        summary.append(f"{name}: {formulas} formulas")
        if name in formula_tabs and literals:
            literal_hits.append(f"{name}: {len(literals)} numeric literals e.g. {literals[:5]}")
    rec("xlsx-structure", "FAIL" if literal_hits else "PASS", "; ".join(summary) + ("; " + "; ".join(literal_hits) if literal_hits else ""))
    rec("xlsx-metadata", "PASS" if wb.properties.creator and wb.properties.creator != "openpyxl" else "NOTE",
        f"creator={wb.properties.creator!r} lastModifiedBy={wb.properties.lastModifiedBy!r}")
    # recalculation: prefer the xlsx skill's recalc.py (LibreOffice macro, rewrites the file in place)
    recalc = glob.glob("/root/.claude/skills/**/xlsx/scripts/recalc.py", recursive=True) or \
        glob.glob(os.path.expanduser("~/.claude/skills/**/xlsx/scripts/recalc.py"), recursive=True)
    if recalc:
        proc = subprocess.run([sys.executable, recalc[0], str(xlsx_path), "180"], capture_output=True, text=True, timeout=400)
        try:
            j = json.loads(proc.stdout)
        except json.JSONDecodeError:
            rec("xlsx-recalc", "FAIL", proc.stdout[-200:] + proc.stderr[-200:])
            return
        ok = j.get("status") == "success" and j.get("total_errors", 1) == 0
        rec("xlsx-recalc", "PASS" if ok else "FAIL", json.dumps(j)[:300])
    else:
        tmp = tempfile.mkdtemp(prefix="recalc_")
        out, log = convert(xlsx_path, "xlsx", tmp)
        if not out:
            rec("xlsx-recalc", "NOTE", f"no recalc.py and LibreOffice conversion failed: {log}")
            return
        wv = openpyxl.load_workbook(out, data_only=True)
        errs = [f"{ws.title}!{c.coordinate}" for ws in wv.worksheets for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("#")]
        rec("xlsx-recalc", "FAIL" if errs else "PASS", f"LibreOffice recalculated a copy; error cells: {errs[:10] or 'none'} (copy at {out})")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("out_dir")
    ap.add_argument("--docx")
    ap.add_argument("--xlsx")
    ap.add_argument("--appendix-heading", default="Appendix")
    ap.add_argument("--min-body-pages", type=int, default=1)
    ap.add_argument("--max-body-pages", type=int, default=4)
    ap.add_argument("--banned-names", default="")
    ap.add_argument("--formula-only-tabs", default="Coverage")
    ap.add_argument("--render", default="")
    ap.add_argument("--json")
    a = ap.parse_args()
    out_dir = Path(a.out_dir)
    check_dashes(out_dir)
    files = client_files(out_dir, a.docx, a.xlsx)
    names = [n.strip() for n in a.banned_names.split(",") if n.strip()]
    if names:
        check_names(files, names)
    check_hype(files)
    if a.docx:
        p = out_dir / a.docx
        check_placeholders(p)
        check_metadata(p)
        render = [int(x) for x in a.render.split(",") if x.strip()]
        check_pages(p, a.appendix_heading, a.min_body_pages, a.max_body_pages, render)
    if a.xlsx:
        check_xlsx(out_dir / a.xlsx, [t.strip() for t in a.formula_only_tabs.split(",")])
    fails = [r for r in results if r["status"] == "FAIL"]
    print(f"\nOVERALL: {'FAIL' if fails else 'PASS'} ({len(fails)} failing checks)")
    if a.json:
        Path(a.json).write_text(json.dumps(results, indent=2))
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
